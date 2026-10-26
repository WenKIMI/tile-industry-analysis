#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 docs/操作指导书.md 转成排版可打印的 DOCX。"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "操作指导书.md"
OUT = ROOT / "docs" / "操作指导书.docx"

BODY_FONT = "等线"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def set_run_font(run, name, size=None, bold=None, color=None, eastasia=True):
    run.font.name = name
    if eastasia:
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(cell_or_par, hex_color):
    el = cell_or_par._element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    el.get_or_add_tcPr().append(shd) if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr().append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(par, text, size=10.5, bold=False, color=None):
    """支持 **粗体** 与 `代码` 两种行内标记。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            # 粗体里可能还嵌着反引号，去掉以免暴露原字符
            r = par.add_run(part[2:-2].replace("`", ""))
            set_run_font(r, BODY_FONT, size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            set_run_font(r, MONO_FONT, size - 1, bold=bold, color=RGBColor(0xC0, 0x39, 0x2B))
        else:
            r = par.add_run(part)
            set_run_font(r, BODY_FONT, size, bold=bold, color=color)


def add_rule(doc):
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BFBFBF")
    pbdr.append(bottom)
    pPr.append(pbdr)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)


def is_table_line(s):
    return s.strip().startswith("|") and s.strip().endswith("|")


def split_row(s):
    return [c.strip() for c in s.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[2:]
    n = len(header)
    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    hr = t.rows[0]
    for i, txt in enumerate(header):
        cell = hr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, txt, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "1F4E79")

    for r_i, row in enumerate(body):
        cells = t.add_row().cells
        for i in range(n):
            txt = row[i] if i < len(row) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_inline(p, txt, size=9)
            if r_i % 2 == 1:
                shade(cells[i], "EDF2F7")

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    md = SRC.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    in_code = False
    code_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # 代码块
        if line.strip().startswith("```"):
            if not in_code:
                in_code, code_buf = True, []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run("\n".join(code_buf))
                set_run_font(r, MONO_FONT, 9)
                shade(p, "F4F4F4")
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        s = line.strip()

        if not s:
            i += 1
            continue

        # 分隔线
        if s == "---":
            add_rule(doc)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2)
            text = text.replace("**", "").replace("`", "")  # 标题内不做行内样式
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                set_run_font(r, HEAD_FONT, 22, bold=True, color=ACCENT)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(14)
            else:
                p = doc.add_paragraph()
                sizes = {2: 15, 3: 12.5, 4: 11}
                r = p.add_run(text)
                set_run_font(r, HEAD_FONT, sizes[level], bold=True, color=ACCENT)
                p.paragraph_format.space_before = Pt(14 if level == 2 else 10)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
            i += 1
            continue

        # 表格
        if is_table_line(s):
            rows = []
            while i < len(lines) and is_table_line(lines[i].strip()):
                rows.append(split_row(lines[i]))
                i += 1
            if len(rows) >= 2:
                add_table(doc, rows)
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, s.lstrip("> ").strip(), size=10, color=RGBColor(0x44, 0x44, 0x44))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # 复选框
        m = re.match(r"^-\s+\[( |x)\]\s*(.*)$", s)
        if m:
            mark = "☑" if m.group(1) == "x" else "☐"
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{mark} {m.group(2)}", size=10.5)
            i += 1
            continue

        # 无序列表
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        add_inline(p, s)
        i += 1

    doc.core_properties.title = "瓷砖行业景气项目 · 操作指导书"
    doc.core_properties.author = ""
    doc.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    build()
