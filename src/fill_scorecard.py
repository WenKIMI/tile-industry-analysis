#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把 data/scorecard_auto.csv 的建议得分回填进 scorecard.xlsx 的「评分卡」Sheet。

只填 5 列，不动公式：
    E 最新值 / F 上期值 / H 得分 / J 打分依据 / K 数据日期
G（同比）和 I（加权分）是公式，保持原样。

用法：
    py -3.13 src/fill_scorecard.py                    # 回填 + 自动备份
    py -3.13 src/fill_scorecard.py --dry-run          # 只看会填什么，不写文件

填完必须人工复核：这里写的是「脚本建议分」，不是最终判断。
不同意的直接改成你认可的分，并把理由写进 J 列。
"""

from __future__ import annotations

import argparse
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "scorecard.xlsx"
AUTO = ROOT / "data" / "scorecard_auto.csv"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(XLSX))
    ap.add_argument("--auto", default=str(AUTO))
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    xlsx, auto = Path(args.xlsx), Path(args.auto)
    if not xlsx.exists():
        raise SystemExit(f"找不到 {xlsx}")
    if not auto.exists():
        raise SystemExit(f"找不到 {auto}，先跑 py -3.13 src/score.py")

    df = pd.read_csv(auto, encoding="utf-8-sig")
    if "建议得分" not in df.columns:
        raise SystemExit("scorecard_auto.csv 缺「建议得分」列")

    wb = load_workbook(xlsx)
    ws = wb["评分卡"]

    # 表头在第 4 行，数据从第 5 行开始到「景气分」行之前
    header_row = 4
    first = header_row + 1
    last = first
    while ws.cell(row=last, column=2).value not in (None, "") and not str(
        ws.cell(row=last, column=2).value
    ).startswith("景气分"):
        last += 1
    last -= 1

    filled, skipped = 0, []
    for r in range(first, last + 1):
        name = ws.cell(row=r, column=2).value
        if not name:
            continue
        hit = df[df["指标"] == name]
        if hit.empty:
            skipped.append(str(name))
            continue

        row = hit.iloc[0]
        latest, prev = row.get("最新值"), row.get("上期值")
        score, why = row.get("建议得分"), row.get("依据")

        # 没有数据的指标不写，保持空白 —— 空白比填 0 更诚实
        if pd.isna(latest):
            skipped.append(str(name))
            continue

        ws.cell(row=r, column=5, value=float(latest))
        if not pd.isna(prev):
            ws.cell(row=r, column=6, value=float(prev))
        ws.cell(row=r, column=8, value=int(score))
        ws.cell(row=r, column=10, value=str(why))
        d = row.get("数据日期")
        ws.cell(row=r, column=11, value=datetime.strptime(str(d), "%Y-%m-%d") if pd.notna(d) else None)
        filled += 1
        print(f"  填 {name:<14} 最新={latest} 上期={prev} 分={int(score)}")

    print(f"\n共填 {filled} 行，跳过 {len(skipped)} 行（无数据，保持空白）")

    if args.dry_run:
        print("\n--dry-run：未写入文件")
        return

    backup = xlsx.with_name(f"{xlsx.stem}_备份{xlsx.suffix}")
    try:
        shutil.copy2(xlsx, backup)
        print(f"已备份到 {backup.name}")
        wb.save(xlsx)
        print(f"已写入 {xlsx.name}")
    except PermissionError:
        raise SystemExit(
            "scorecard.xlsx 正被 Excel/WPS 打开，写不进去。关掉后重跑本脚本。"
        )


if __name__ == "__main__":
    main()
